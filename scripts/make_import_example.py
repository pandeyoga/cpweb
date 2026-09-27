#!/usr/bin/env python
"""scripts/make_import_example.py — Generator FILE CONTOH impor produk (XLSX + CSV).

Berbeda dengan endpoint `/api/admin/products/io/template` (butuh login admin), file hasil
script ini disimpan di `frontend/public/templates/` sehingga bisa diunduh langsung dari
storefront tanpa autentikasi:

    <REACT_APP_BACKEND_URL>/templates/contoh-impor-produk-collector-parfum.xlsx

Isi file:
  • Sheet "Produk"    : header + 9 baris contoh (produk 2D, 3D, dan 4D) — SIAP DIEDIT.
  • Sheet "Panduan"   : aturan pengisian (diambil dari services.product_bulk.INSTRUCTIONS).
  • Sheet "Referensi" : daftar slug kategori / occasion / character yang VALID di database.

Jalankan: python scripts/make_import_example.py
"""
import asyncio
import csv
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from dotenv import load_dotenv  # noqa: E402

load_dotenv(ROOT / "backend" / ".env")

from motor.motor_asyncio import AsyncIOMotorClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from services.product_bulk import INSTRUCTIONS  # noqa: E402
from services.product_io import EXPORT_COLUMNS  # noqa: E402

OUT_DIR = ROOT / "frontend" / "public" / "templates"
BASENAME = "contoh-impor-produk-collector-parfum"

MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")


def _row(base, dims, price, stock, compare="", sku=""):
    """Bangun satu baris varian: base produk + pasangan dimensi option{i}_name/value."""
    row = {k: "" for k in EXPORT_COLUMNS}
    row.update(base)
    for i, (name, value) in enumerate(dims[:4], start=1):
        row[f"option{i}_name"] = name
        row[f"option{i}_value"] = value
    row["variant_price"] = price
    row["variant_stock"] = stock
    row["variant_compare_at_price"] = compare
    row["variant_sku"] = sku
    return row


def example_rows():
    # --- Produk 1: 2 dimensi (Konsentrasi × Ukuran) ---
    p1 = {
        "slug": "contoh-mawar-pagi",
        "name": "Contoh Mawar Pagi",
        "brand": "Collector",
        "category": "floral",
        "concentration": "EDT",
        "gender": "Wanita",
        "description": "Mawar segar dengan sentuhan sitrus — ringan untuk pemakaian harian.",
        "ingredients": "Alcohol Denat., Parfum (Fragrance), Aqua",
        "tags": "Rose, Fresh",
        "occasions": "daily-wear, office",
        "characters": "floral, fresh",
        "best_seller": "false",
        "is_new": "true",
        "status": "active",
        "images": "",
        "video_url": "",
    }
    # --- Produk 2: 3 dimensi (Konsentrasi × Tipe × Ukuran) ---
    p2 = {
        "slug": "contoh-oud-malam",
        "name": "Contoh Oud Malam",
        "brand": "Collector",
        "category": "amber",
        "concentration": "EDP",
        "gender": "Unisex",
        "description": "Oud hangat berbalut amber — karakter kuat untuk acara malam.",
        "ingredients": "Alcohol Denat., Parfum (Fragrance), Aqua",
        "tags": "Oud, Amber, Woody",
        "occasions": "evening, date-night, party",
        "characters": "amber-oud, woody, musky",
        "best_seller": "true",
        "is_new": "false",
        "status": "active",
        "images": "",
        "video_url": "",
    }
    # --- Produk 3: 4 dimensi (Konsentrasi × Tipe × Ukuran × Kemasan) ---
    p3 = {
        "slug": "contoh-kayu-hujan",
        "name": "Contoh Kayu Hujan",
        "brand": "Collector",
        "category": "woody",
        "concentration": "EDP",
        "gender": "Pria",
        "description": "Cendana sesudah hujan — tenang, bersih, dan maskulin.",
        "ingredients": "Alcohol Denat., Parfum (Fragrance), Aqua",
        "tags": "Sandalwood, Petrichor",
        "occasions": "office, daily-wear",
        "characters": "woody, green",
        "best_seller": "false",
        "is_new": "false",
        "status": "active",
        "images": "",
        "video_url": "",
    }
    return [
        # 2 dimensi — 3 varian ukuran
        _row(p1, [("Konsentrasi", "EDT"), ("Ukuran", "30ml")], 150000, 20),
        _row(p1, [("Konsentrasi", "EDT"), ("Ukuran", "50ml")], 250000, 12, compare=299000),
        _row(p1, [("Konsentrasi", "EDT"), ("Ukuran", "100ml")], 425000, 6),
        # 3 dimensi — matriks Tipe × Ukuran
        _row(p2, [("Konsentrasi", "EDP"), ("Tipe", "Standard"), ("Ukuran", "50ml")], 450000, 10),
        _row(p2, [("Konsentrasi", "EDP"), ("Tipe", "Standard"), ("Ukuran", "100ml")], 780000, 5),
        _row(p2, [("Konsentrasi", "EDP"), ("Tipe", "Premium"), ("Ukuran", "50ml")], 620000, 4),
        _row(p2, [("Konsentrasi", "EDP"), ("Tipe", "Premium"), ("Ukuran", "100ml")], 990000, 2,
             compare=1250000),
        # 4 dimensi — tambah dimensi Kemasan
        _row(p3, [("Konsentrasi", "EDP"), ("Tipe", "Standard"), ("Ukuran", "50ml"),
                  ("Kemasan", "Botol")], 395000, 8),
        _row(p3, [("Konsentrasi", "EDP"), ("Tipe", "Standard"), ("Ukuran", "50ml"),
                  ("Kemasan", "Gift Box")], 465000, 3, sku="KAYUHUJAN-STD-50ML-GIFT"),
    ]


EXTRA_GUIDE = [
    "",
    "KOLOM TAKSONOMI (opsional, tapi sangat disarankan)",
    "• 'occasions' dan 'characters' = daftar SLUG dipisah koma (lihat sheet 'Referensi').",
    "  Dipakai untuk filter Occasion & Character di halaman Toko.",
    "• Bila kolom ini DIKOSONGKAN saat impor mode Upsert, nilai lama produk DIPERTAHANKAN",
    "  (tidak terhapus). Untuk mengosongkan, ubah lewat halaman Admin > Produk.",
    "",
    "HARGA & STOK",
    "• 'variant_price' dan 'variant_compare_at_price' = ANGKA BULAT rupiah tanpa titik/Rp",
    "  (contoh: 250000). 'variant_stock' = angka bulat >= 0.",
    "• 'variant_compare_at_price' (harga coret) harus LEBIH BESAR dari variant_price.",
    "",
    "GAMBAR",
    "• 'images' = daftar URL dipisah koma. Bila kosong, sistem memakai artwork botol otomatis.",
    "  Untuk foto asli, unggah dulu di Admin > Konten (media) lalu tempel URL-nya di sini.",
    "",
    "CARA IMPOR",
    "1. Isi/ubah sheet 'Produk' (JANGAN ubah nama kolom baris pertama).",
    "2. Simpan sebagai .xlsx atau .csv.",
    "3. Buka Admin > Impor / Ekspor, unggah file, cek pemetaan kolom + pratinjau.",
    "4. Pilih mode: 'Tambah saja (aman)' untuk produk baru, atau 'Upsert' untuk memperbarui",
    "   produk yang sudah ada (dicocokkan lewat slug/nama).",
]


async def fetch_reference():
    client = AsyncIOMotorClient(MONGO_URL)
    db = client[DB_NAME]
    cats = await db.categories.find({}, {"slug": 1, "name": 1}).to_list(200)
    occs = await db.occasions.find({}, {"slug": 1, "name": 1}).to_list(200)
    chars = await db.characters.find({}, {"slug": 1, "name": 1}).to_list(200)
    client.close()
    return cats, occs, chars


def build_xlsx(rows, cats, occs, chars) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Produk"
    ws.append(EXPORT_COLUMNS)
    for r in rows:
        ws.append([r.get(k, "") for k in EXPORT_COLUMNS])

    head_fill = PatternFill("solid", fgColor="141414")
    head_font = Font(color="FFFCF6", bold=True, size=10)
    for idx, col in enumerate(EXPORT_COLUMNS, start=1):
        c = ws.cell(row=1, column=idx)
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        width = 34 if col in ("description", "ingredients") else max(12, min(24, len(col) + 4))
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("Panduan")
    for line in list(INSTRUCTIONS) + EXTRA_GUIDE:
        guide.append([line])
    guide.column_dimensions["A"].width = 100
    guide["A1"].font = Font(bold=True, size=12)

    ref = wb.create_sheet("Referensi")
    ref.append(["JENIS", "SLUG (dipakai di file impor)", "NAMA TAMPILAN"])
    for c in ("A", "B", "C"):
        ref[f"{c}1"].font = Font(bold=True)
        ref.column_dimensions[c].width = 32
    for c in sorted(cats, key=lambda x: x.get("slug", "")):
        ref.append(["category", c.get("slug", ""), c.get("name", "")])
    for o in sorted(occs, key=lambda x: x.get("slug", "")):
        ref.append(["occasions", o.get("slug", ""), o.get("name", "")])
    for ch in sorted(chars, key=lambda x: x.get("slug", "")):
        ref.append(["characters", ch.get("slug", ""), ch.get("name", "")])
    ref.append(["gender", "Pria / Wanita / Unisex", "pilih salah satu"])
    ref.append(["concentration", "EDP / EDT", "pilih salah satu"])
    ref.append(["status", "active / archived", "archived = disembunyikan dari storefront"])
    ref.append(["best_seller / is_new", "true / false", "boleh juga 1 / 0, ya / tidak"])

    import io
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_csv(rows) -> bytes:
    import io
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=EXPORT_COLUMNS, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow(r)
    return buf.getvalue().encode("utf-8-sig")  # BOM -> Excel ramah


async def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    cats, occs, chars = await fetch_reference()
    rows = example_rows()

    xlsx_path = OUT_DIR / f"{BASENAME}.xlsx"
    csv_path = OUT_DIR / f"{BASENAME}.csv"
    xlsx_path.write_bytes(build_xlsx(rows, cats, occs, chars))
    csv_path.write_bytes(build_csv(rows))

    print(f"✓ {xlsx_path}  ({xlsx_path.stat().st_size:,} bytes)")
    print(f"✓ {csv_path}  ({csv_path.stat().st_size:,} bytes)")
    print(f"  kolom : {len(EXPORT_COLUMNS)} -> {', '.join(EXPORT_COLUMNS)}")
    print(f"  contoh: {len(rows)} baris varian / 3 produk (2D, 3D, 4D)")
    print(f"  ref   : {len(cats)} kategori, {len(occs)} occasion, {len(chars)} character")
    return 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
