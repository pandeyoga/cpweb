#!/usr/bin/env python
"""scripts/verify_user_import.py — VERIFIKASI impor katalog klien (file XLSX asli user).

Tujuan: membuktikan (bukan mengklaim) bahwa file
`tests/user_uploads/IMPOR_PRODUK_COLLECTOR_PARFUM_1.xlsx` benar-benar bisa diimpor
lewat endpoint HTTP produksi `/api/admin/products/io/*`, dan mendokumentasikan
apa yang terjadi pada baris berharga 0.

Jalankan: python scripts/verify_user_import.py
"""
import os
import sys
from collections import OrderedDict

import requests
from openpyxl import load_workbook

BASE = os.environ.get("BACKEND_BASE", "http://localhost:8001/api")
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "admin@collectorparfum.id")
ADMIN_PASS = os.environ.get("ADMIN_PASS", "Admin#2026")
XLSX = "tests/user_uploads/IMPOR_PRODUK_COLLECTOR_PARFUM_1.xlsx"

G, R, Y, C, X, B = "\033[92m", "\033[91m", "\033[93m", "\033[96m", "\033[0m", "\033[1m"
PASS = FAIL = 0


def ok(label, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  {G}[PASS]{X} {label}" + (f"  {extra}" if extra else ""))
    else:
        FAIL += 1
        print(f"  {R}[FAIL]{X} {label}" + (f"  {extra}" if extra else ""))


def head(t):
    print(f"\n{C}{B}--- {t} ---{X}")


def read_sheet(path, sheet="Produk"):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = []
    for r in it:
        if not any(c not in (None, "") for c in r):
            continue
        rows.append({hdr[i]: ("" if v is None else str(v).strip())
                     for i, v in enumerate(r) if i < len(hdr)})
    wb.close()
    return hdr, rows


def group(rows):
    g = OrderedDict()
    for r in rows:
        g.setdefault((r.get("slug") or r.get("name") or "").strip(), []).append(r)
    return g


def login():
    r = requests.post(f"{BASE}/auth/login",
                      json={"email": ADMIN_EMAIL, "password": ADMIN_PASS}, timeout=30)
    r.raise_for_status()
    return {"Authorization": f"Bearer {r.json()['token']}"}


def purge(slugs):
    """Hapus permanen produk uji langsung dari Mongo (DELETE API = soft-delete)."""
    import asyncio

    from dotenv import load_dotenv
    from motor.motor_asyncio import AsyncIOMotorClient
    load_dotenv("backend/.env")

    async def run():
        c = AsyncIOMotorClient(os.environ["MONGO_URL"])
        db = c[os.environ["DB_NAME"]]
        res = await db.products.delete_many({"slug": {"$in": list(slugs)}})
        return res.deleted_count
    return asyncio.run(run())


def main():
    if not os.path.exists(XLSX):
        print(f"{R}File tidak ditemukan: {XLSX}{X}")
        return 1

    print(f"{C}{B}=========================================================={X}")
    print(f"{C}{B}  VERIFIKASI IMPOR KATALOG KLIEN — {XLSX}{X}")
    print(f"{C}{B}=========================================================={X}")

    hdr, rows = read_sheet(XLSX)
    groups = group(rows)

    head("1. Struktur file")
    ok("sheet 'Produk' terbaca", len(rows) > 0, f"{len(rows)} baris data")
    ok("28 kolom header terbaca", len(hdr) == 28, f"{len(hdr)} kolom")
    ok("jumlah produk unik terdeteksi", len(groups) == 1071, f"{len(groups)} produk")
    ok("setiap produk 6 varian (3 Ukuran x 2 Tipe)",
       all(len(v) == 6 for v in groups.values()))
    ok("dimensi = Ukuran + Tipe (2D)",
       all(r["option1_name"] == "Ukuran" and r["option2_name"] == "Tipe" for r in rows))
    zero_price = sum(1 for r in rows if float(r["variant_price"] or 0) <= 0)
    ok("SEMUA harga di file = 0 (placeholder klien)", zero_price == len(rows),
       f"{zero_price}/{len(rows)}")

    head("2. Referensi taksonomi vs DB (anti FK menggantung)")
    cats = {c["slug"] for c in requests.get(f"{BASE}/categories", timeout=30).json()}
    occs = {c["slug"] for c in requests.get(f"{BASE}/occasions", timeout=30).json()}
    chs = {c["slug"] for c in requests.get(f"{BASE}/characters", timeout=30).json()}
    file_cats = {r["category"] for r in rows if r["category"]}
    file_occ = {o.strip() for r in rows for o in (r["occasions"] or "").split(",") if o.strip()}
    file_ch = {o.strip() for r in rows for o in (r["characters"] or "").split(",") if o.strip()}
    ok("semua slug category ada di DB", file_cats <= cats, f"beda={sorted(file_cats - cats)}")
    ok("semua slug occasions ada di DB", file_occ <= occs, f"beda={sorted(file_occ - occs)}")
    ok("semua slug characters ada di DB", file_ch <= chs, f"beda={sorted(file_ch - chs)}")

    H = login()
    mapping = {h: h for h in hdr if h}

    # ---- sampel 3 produk pertama (18 baris) ----
    sample_keys = list(groups)[:3]
    sample_raw = [r for k in sample_keys for r in groups[k]]
    print(f"\n{Y}sampel uji: {sample_keys}{X}")

    head("3. Impor APA ADANYA (harga 0) — harus DITOLAK dengan pesan jelas, bukan 5xx")
    r = requests.post(f"{BASE}/admin/products/io/validate", headers=H, timeout=120,
                      json={"rows": sample_raw, "mapping": mapping})
    ok("POST /validate tidak 5xx", r.status_code < 500, f"HTTP {r.status_code}")
    if r.status_code == 200:
        d = r.json()
        reps = d.get("reports") or []
        n_err = sum(1 for row in reps if row.get("errors"))
        ok("baris harga 0 ditandai error (tidak lolos diam-diam)", n_err == len(sample_raw),
           f"{n_err}/{len(sample_raw)} baris error")
        msgs = {e for row in reps for e in (row.get("errors") or [])}
        ok("pesan error menyebut harga", any("harga" in m.lower() or "price" in m.lower()
                                            for m in msgs), f"{sorted(msgs)[:2]}")
        ok("summary.error = 18, products = 0", d.get("summary", {}).get("error") == 18
           and d.get("summary", {}).get("products") == 0, f"{d.get('summary')}")
    rc = requests.post(f"{BASE}/admin/products/io/commit", headers=H, timeout=300,
                       json={"rows": sample_raw, "mapping": mapping, "mode": "add-only"})
    ok("POST /commit tidak 5xx", rc.status_code < 500, f"HTTP {rc.status_code}")
    if rc.status_code == 200:
        d = rc.json()
        ok("commit harga 0 -> 0 produk dibuat (data kotor ditahan)",
           int(d.get("created") or 0) == 0, f"created={d.get('created')} failed={d.get('failed')}")
        ok("commit melaporkan row_errors ke admin (bukan sunyi)",
           int(d.get("row_errors") or 0) == len(sample_raw), f"row_errors={d.get('row_errors')}")

    head("4. Impor dengan harga placeholder (jalur scripts/import_user_catalog.py)")
    norm = []
    for r0 in sample_raw:
        r1 = dict(r0)
        r1["variant_price"] = "1000"
        r1["variant_stock"] = "0"
        r1["concentration"] = "EDP"
        r1["status"] = "archived"
        norm.append(r1)
    rv = requests.post(f"{BASE}/admin/products/io/validate", headers=H, timeout=120,
                       json={"rows": norm, "mapping": mapping})
    ok("validate setelah normalisasi = 200", rv.status_code == 200, f"HTTP {rv.status_code}")
    if rv.status_code == 200:
        d = rv.json()
        bad = [row for row in (d.get("reports") or []) if row.get("errors")]
        ok("0 baris error setelah normalisasi", not bad,
           f"contoh={bad[0].get('errors') if bad else ''}")
        prods = d.get("products") or []
        ok("grouping -> 3 produk", len(prods) == 3, f"{len(prods)} produk")
        if prods:
            p0 = prods[0]
            dims = p0.get("options") or []
            ok("produk hasil grouping punya 2 dimensi", len(dims) == 2,
               f"{[o.get('name') for o in dims]}")
            ok("produk hasil grouping punya 6 varian", len(p0.get("variants") or []) == 6,
               f"{len(p0.get('variants') or [])} varian")

    rc = requests.post(f"{BASE}/admin/products/io/commit", headers=H, timeout=300,
                       json={"rows": norm, "mapping": mapping, "mode": "add-only"})
    ok("commit = 200", rc.status_code == 200, f"HTTP {rc.status_code}")
    created = 0
    if rc.status_code == 200:
        d = rc.json()
        created = int(d.get("created") or 0)
        ok("3 produk dibuat", created == 3,
           f"created={created} updated={d.get('updated')} skipped={d.get('skipped')} failed={d.get('failed')}")
        if d.get("errors"):
            print(f"    {Y}errors: {d['errors'][:3]}{X}")

    head("5. Verifikasi hasil di DB & storefront")
    ap = requests.get(f"{BASE}/admin/products", headers=H, timeout=60,
                      params={"q": "1000 Bunga", "limit": 20})
    ok("GET /admin/products menemukan produk impor", ap.status_code == 200)
    found = None
    if ap.status_code == 200:
        items = ap.json().get("items") if isinstance(ap.json(), dict) else ap.json()
        found = next((p for p in (items or []) if p.get("slug") == sample_keys[0]), None)
        ok("produk pertama ada di daftar admin", found is not None, f"slug={sample_keys[0]}")
    if found:
        ok("status = archived (tersembunyi dari storefront)",
           found.get("status") == "archived", f"status={found.get('status')}")
        variants = found.get("variants") or []
        ok("6 varian tersimpan", len(variants) == 6, f"{len(variants)} varian")
        ok("SKU auto-generate untuk semua varian",
           all(v.get("sku") for v in variants),
           f"contoh={variants[0].get('sku') if variants else ''}")
        opts = found.get("options") or []
        ok("options = ['Ukuran','Tipe']", [o.get("name") for o in opts] == ["Ukuran", "Tipe"],
           f"{[o.get('name') for o in opts]}")
        vols = found.get("volumes") or []
        ok("volumes[] turunan: ml > 0 untuk semua varian",
           len(vols) == 6 and all(int(v.get("ml") or 0) > 0 for v in vols),
           f"ml={[v.get('ml') for v in vols]}")
        ok("volumes[].type = gabungan dimensi non-ukuran",
           all(v.get("type") for v in vols), f"type={[v.get('type') for v in vols][:3]}")
        ok("occasions tersimpan", len(found.get("occasions") or []) == 3,
           f"{found.get('occasions')}")
        ok("characters tersimpan", len(found.get("characters") or []) >= 1,
           f"{found.get('characters')}")
        ok("brand tersimpan", bool(found.get("brand")), f"brand={found.get('brand')}")
        ok("harga placeholder 1000 tersimpan",
           all(int(v.get("price") or 0) == 1000 for v in variants))

    sf = requests.get(f"{BASE}/products/{sample_keys[0]}", timeout=30)
    ok("produk archived TIDAK tampil di storefront (404)", sf.status_code == 404,
       f"HTTP {sf.status_code}")

    head("6. Aktivasi: ubah status -> active, produk muncul di storefront")
    if found:
        up = requests.patch(f"{BASE}/admin/products/{found['id']}", headers=H, timeout=60,
                           json={"status": "active"})
        if up.status_code == 405 or up.status_code == 404:
            # fallback: PUT full
            full = dict(found)
            full["status"] = "active"
            up = requests.put(f"{BASE}/admin/products/{found['id']}", headers=H,
                              timeout=60, json=full)
        ok("update status -> active berhasil", up.status_code == 200, f"HTTP {up.status_code}")
        sf2 = requests.get(f"{BASE}/products/{sample_keys[0]}", timeout=30)
        ok("produk aktif tampil di storefront (200)", sf2.status_code == 200,
           f"HTTP {sf2.status_code}")
        if sf2.status_code == 200:
            p = sf2.json()
            ok("PDP membawa 6 varian", len(p.get("variants") or []) == 6)
            ok("price_min/price_max terisi (flattened price_range)",
               int(p.get("price_min") or 0) > 0 and int(p.get("price_max") or 0) > 0,
               f"min={p.get('price_min')} max={p.get('price_max')} price={p.get('price')}")

    head("7. Upsert memperbarui harga tanpa duplikasi")
    norm2 = []
    for r0 in norm:
        r1 = dict(r0)
        r1["variant_price"] = "250000"
        r1["status"] = "active"
        norm2.append(r1)
    ru = requests.post(f"{BASE}/admin/products/io/commit", headers=H, timeout=300,
                       json={"rows": norm2, "mapping": mapping, "mode": "upsert"})
    ok("commit upsert = 200", ru.status_code == 200, f"HTTP {ru.status_code}")
    if ru.status_code == 200:
        d = ru.json()
        ok("3 produk diperbarui, 0 dibuat", int(d.get("updated") or 0) == 3
           and int(d.get("created") or 0) == 0,
           f"created={d.get('created')} updated={d.get('updated')}")
    sf3 = requests.get(f"{BASE}/products/{sample_keys[0]}", timeout=30)
    if sf3.status_code == 200:
        vs = sf3.json().get("variants") or []
        ok("harga varian jadi 250.000", all(int(v.get("price") or 0) == 250000 for v in vs),
           f"{[v.get('price') for v in vs][:3]}")
        ok("tetap 6 varian (tidak duplikat)", len(vs) == 6, f"{len(vs)} varian")
        ok("occasions TETAP utuh setelah upsert (anti data-loss)",
           len(sf3.json().get("occasions") or []) == 3, f"{sf3.json().get('occasions')}")

    head("8. Bersihkan produk uji")
    n = purge(sample_keys)
    ok("produk uji dihapus permanen dari DB", n == 3, f"{n} dihapus")

    print(f"\n{C}{B}=========================================================={X}")
    print(f"  {G}PASS: {PASS}{X}    {R}FAIL: {FAIL}{X}")
    print(f"{C}{B}=========================================================={X}")
    if FAIL == 0:
        print(f"{G}{B}  SEMUA VERIFIKASI IMPOR LULUS.{X}\n")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
