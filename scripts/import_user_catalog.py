#!/usr/bin/env python
"""scripts/import_user_catalog.py — jalur impor ALTERNATIF (bertahap via API) untuk file besar.

Diaudit ulang (butir 21):
  - TIDAK mengarang harga: harga 0/kosong dikirim apa adanya; produk yang punya varian tanpa harga
    sah dipaksa `status=draft` (publikasi ditahan) — isi harga lewat wizard/`--prices <csv>` resmi pemilik.
  - TIDAK memulihkan `concentration`/`ingredients` (kolom dibuang bila ada).
  - Mempertahankan `tier`, `date_night`, `characters` apa adanya dari file.
  - TIDAK ada password bawaan: ADMIN_EMAIL, ADMIN_PASS, BACKEND_BASE WAJIB dari environment.
  - Batch HTTP gagal / baris gagal = GAGAL: exit code 1 + ringkasan JSON (`--summary`) berisi batch gagal;
    jalankan ulang dengan `--resume <summary.json>` untuk mengulang hanya batch yang gagal (idempoten, mode upsert).

Pakai:
  BACKEND_BASE=https://domain/api ADMIN_EMAIL=... ADMIN_PASS=... \\
  python scripts/import_user_catalog.py <file.xlsx> [--mode upsert] [--batch 100] [--summary out.json] [--resume out.json]
"""
import argparse
import json
import os
import sys
from collections import OrderedDict

import requests
from openpyxl import load_workbook

DROP_COLS = {"concentration", "ingredients"}


def read_sheet(path, sheet="Produk"):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    out = []
    for r in rows:
        if any(c not in (None, "") for c in r):
            out.append({header[i]: ("" if v is None else str(v).strip())
                        for i, v in enumerate(r) if i < len(header) and header[i] not in DROP_COLS})
    wb.close()
    return [h for h in header if h and h not in DROP_COLS], out


def _price_ok(v):
    try:
        return float(v) > 0
    except (TypeError, ValueError):
        return False


def group_and_hold(rows):
    """Kelompokkan per slug; produk dengan varian tanpa harga sah → draft (tahan publikasi)."""
    groups = OrderedDict()
    for r in rows:
        groups.setdefault((r.get("slug") or r.get("name") or "").strip(), []).append(r)
    held = 0
    for rs in groups.values():
        if not all(_price_ok(r.get("variant_price")) for r in rs):
            held += 1
            for r in rs:
                r["status"] = "draft"
    return groups, held


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--mode", default="upsert", choices=["add-only", "upsert"])
    ap.add_argument("--batch", type=int, default=100, help="jumlah PRODUK per request")
    ap.add_argument("--summary", default="import_summary.json")
    ap.add_argument("--resume", default=None, help="ulang hanya batch gagal dari ringkasan sebelumnya")
    a = ap.parse_args()
    base, email, password = (os.environ.get(k) for k in ("BACKEND_BASE", "ADMIN_EMAIL", "ADMIN_PASS"))
    if not (base and email and password):
        print("BACKEND_BASE, ADMIN_EMAIL, ADMIN_PASS wajib diisi di environment (tanpa nilai bawaan).")
        return 2

    header, rows = read_sheet(a.path)
    groups, held = group_and_hold(rows)
    keys = list(groups)
    batches = list(range(0, len(keys), a.batch))
    if a.resume:
        prev = json.load(open(a.resume))
        batches = [b * a.batch for b in prev.get("failed_batches", [])]
    print(f"{len(keys)} produk, {len(batches)} batch, {held} produk ditahan (draft) karena harga belum sah")

    tok = requests.post(f"{base.rstrip('/')}/auth/login", json={"email": email, "password": password}, timeout=30)
    if tok.status_code != 200:
        print(f"Login gagal: HTTP {tok.status_code}")
        return 1
    hdr = {"Authorization": f"Bearer {tok.json()['token']}"}
    summary = {"file": a.path, "mode": a.mode, "batch": a.batch, "held_draft": held,
               "created": 0, "updated": 0, "skipped": 0, "failed": 0, "failed_batches": [], "errors": []}
    for start in batches:
        n = start // a.batch
        chunk = [r for k in keys[start:start + a.batch] for r in groups[k]]
        try:
            res = requests.post(f"{base.rstrip('/')}/admin/products/io/commit", headers=hdr, timeout=600,
                                json={"rows": chunk, "mapping": {h: h for h in header}, "mode": a.mode})
            d = res.json() if res.status_code == 200 else {}
        except Exception as e:  # jaringan → batch gagal (bukan sukses diam-diam)
            res, d = None, {"errors": [str(e)[:200]]}
        failed = res is None or res.status_code != 200 or int(d.get("failed") or 0) > 0
        for k in ("created", "updated", "skipped", "failed"):
            summary[k] += int(d.get(k) or 0)
        if failed:
            summary["failed_batches"].append(n)
            summary["errors"].extend((d.get("errors") or [f"HTTP {getattr(res, 'status_code', '-')}"])[:5])
        print(f"batch {n + 1}: {'GAGAL' if failed else 'ok'} {d.get('created', 0)} dibuat, {d.get('updated', 0)} diperbarui")
    json.dump(summary, open(a.summary, "w"), ensure_ascii=False, indent=2)
    print(f"Ringkasan → {a.summary}: gagal {len(summary['failed_batches'])} batch")
    return 1 if summary["failed_batches"] else 0


if __name__ == "__main__":
    sys.exit(main())
